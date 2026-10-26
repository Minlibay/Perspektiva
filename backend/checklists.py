"""
Реестр типов чек-листов и универсальный разбор / заполнение их шаблонов (.docx/.docm).

Колонки таблицы чек-листа находятся по заголовкам («Область проверки», «ОК», «NOK»,
«Проблемные зоны»), а не по фиксированным номерам: новые редакции шаблонов добавляют
и переставляют колонки (План АУДИТА 2026 — 7 колонок с пустой первой). Правила проверки
привязываются к тексту пункта (см. _plan_rule_id / _svod_rule_id в main.py).
Спецификация всех чек-листов — docs/ЧЕК-ЛИСТЫ.md.
"""

import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt


# Порядок важен: специфичные заголовки раньше общих («Заключения – 1 этап» раньше
# «Заключения»). signatures сравниваются с заголовком без пробелов и знаков (_squash).
# supported — есть ли пайплайн проверки; остальные распознаются, но не обрабатываются
# (этапы внедрения — docs/ЧЕК-ЛИСТЫ.md, §7).
CHECKLIST_TYPES = [
    {"key": "zakl1", "title": "Заключение — 1 этап", "checked_doc": "zakl1",
     "signatures": ["проверкизаключения1этап"], "supported": True, "stage": 3},
    {"key": "zakl", "title": "Заключение", "checked_doc": "zakl",
     "signatures": ["проверкизаключения"], "supported": True, "stage": 2},
    {"key": "otchet1", "title": "Отчет 1 этапа", "checked_doc": "otchet1",
     "signatures": ["проверкиотчета1этапа"], "supported": True, "stage": 3},
    {"key": "akt", "title": "Акт", "checked_doc": "akt",
     "signatures": ["проверкиактпорезультатам"], "supported": True, "stage": 2},
    {"key": "loo_org", "title": "Лист оценки организацией", "checked_doc": "loo_org",
     "signatures": ["проверкилистаоценкиорганизац"], "supported": True, "stage": 1},
    {"key": "loo_exp", "title": "Лист оценки экспертами", "checked_doc": "loo_exp",
     "signatures": ["проверкилистаоценкиэксперт"], "supported": True, "stage": 1},
    {"key": "otchet_kd", "title": "Отчет КД", "checked_doc": "otchet_kd",
     "signatures": ["проверкиотчеткд"], "supported": True, "stage": 1},
    {"key": "pkd", "title": "ПКД", "checked_doc": "pkd",
     "signatures": ["проверкипкд"], "supported": True, "stage": 1},
    {"key": "pn", "title": "Протоколы несоответствий", "checked_doc": "pn",
     "signatures": ["проверкипротоколоврегистрациинесоответстви"], "supported": True, "stage": 1},
    # План и Сводный акт распознаются по структуре (заголовка «Чек лист проверки …» у них нет).
    {"key": "plan", "title": "План АУДИТА", "checked_doc": "plan",
     "signatures": [], "supported": True, "stage": 0},
    {"key": "svod", "title": "Сводный акт", "checked_doc": "svod",
     "signatures": [], "supported": True, "stage": 0},
]

CHECKLIST_BY_KEY = {c["key"]: c for c in CHECKLIST_TYPES}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").lower().replace("ё", "е")).strip()


def _squash(s: str) -> str:
    """Только буквы и цифры: «Чек лист проверки Заключения – 1этап» → «чеклистпроверкизаключения1этап»."""
    return re.sub(r"[^0-9a-zа-я]", "", (s or "").lower().replace("ё", "е"))


def open_docx(file_path):
    """python-docx не открывает .docm (другой content-type main part).
    Перепаковываем в памяти: меняем macroEnabled.main → document.main."""
    if Path(str(file_path)).suffix.lower() != ".docm":
        return Document(str(file_path))
    src_macro = b"application/vnd.ms-word.document.macroEnabled.main+xml"
    dst_main = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
    buf = BytesIO()
    with zipfile.ZipFile(str(file_path), "r") as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(src_macro, dst_main)
            zout.writestr(item, data)
    buf.seek(0)
    return Document(buf)


def _as_document(doc_or_path):
    return doc_or_path if hasattr(doc_or_path, "tables") else open_docx(doc_or_path)


def _unique_row_texts(row) -> list:
    """Тексты ячеек строки без повторов объединённых ячеек."""
    seen, texts = set(), []
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        texts.append(c.text.strip())
    return texts


def _is_title_row(row) -> bool:
    """Строка-титул: одна ячейка, объединённая на всю ширину таблицы."""
    return len({id(c._tc) for c in row.cells}) == 1


def _is_svod_structure(doc) -> bool:
    for table in doc.tables:
        if not table.rows or len(table.rows[0].cells) < 6:
            continue
        header = [c.text.strip() for c in table.rows[0].cells]
        first_rows = " ".join(c.text.strip().upper() for row in table.rows[:3] for c in row.cells)
        if any("Проблемные зоны" in c for c in header) and "СВОДНЫЙ АКТ" in first_rows:
            return True
    return False


def detect_checklist_type(doc_or_path) -> str:
    """Тип чек-листа: ключ из CHECKLIST_TYPES или 'unknown'.

    Новые чек-листы опознаются по заголовку «Чек лист проверки …» (абзац или
    строка-титул таблицы). План АУДИТА и Сводный акт заголовка не имеют — для них
    структурные признаки: таблица шапки «Наименования Заявителя» / таблица с колонкой
    «Область проверки»; у Сводного — «Проблемные зоны» и титул «СВОДНЫЙ АКТ».
    """
    try:
        doc = _as_document(doc_or_path)
    except Exception as e:
        print(f"[checklists] не удалось открыть {doc_or_path}: {e}")
        return "unknown"

    title_parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        if table.rows:
            title_parts.append(" ".join(_unique_row_texts(table.rows[0])))
    title = _squash(" ".join(title_parts))
    for ctype in CHECKLIST_TYPES:
        if any(sig in title for sig in ctype["signatures"]):
            return ctype["key"]

    if _is_svod_structure(doc):
        return "svod"
    for table in doc.tables:
        for row in table.rows[:4]:
            texts = [c.text.strip() for c in row.cells]
            if any("Наименования Заявителя" in t or "Область проверки" in t for t in texts):
                return "plan"
    return "unknown"


def _find_layout(table) -> Optional[dict]:
    """Колонки таблицы чек-листа по заголовкам.

    Возвращает {header_row, area, comment, ok, nok, problems, docs, num} (индексы
    ячеек row.cells) или None, если таблица не похожа на чек-лист.
    Шаблоны без строки заголовка (Чек лист Акт: титул + 5 колонок) и с пустыми
    заголовками ОК/NOK (План АУДИТА v1) получают раскладку
    [Область | Комментарий | ОК | NOK | Источник] от колонки «Область проверки».
    """
    if not table.rows:
        return None
    ncols = max(len(r.cells) for r in table.rows)
    if ncols < 5:
        return None

    layout = {"header_row": None, "area": None, "comment": None, "ok": None, "nok": None,
              "problems": None, "docs": None, "num": None}
    for ri, row in enumerate(table.rows[:4]):
        found = False
        for ci, cell in enumerate(row.cells):
            t = _norm(cell.text)
            if not t:
                continue
            if t in ("ок", "ok"):
                layout["ok"] = ci if layout["ok"] is None else layout["ok"]
                found = True
            elif t in ("nok", "нок"):
                layout["nok"] = ci if layout["nok"] is None else layout["nok"]
                found = True
            elif "область проверки" in t or t in ("наименование", "наимнование"):
                layout["area"] = ci if layout["area"] is None else layout["area"]
                found = True
            elif t.startswith("комментари"):
                layout["comment"] = ci if layout["comment"] is None else layout["comment"]
                found = True
            elif "проблемные зоны" in t:
                layout["problems"] = ci if layout["problems"] is None else layout["problems"]
                found = True
            elif t.startswith("доки"):
                layout["docs"] = ci if layout["docs"] is None else layout["docs"]
                found = True
            elif t == "№":
                layout["num"] = ci if layout["num"] is None else layout["num"]
                found = True
        if found:
            layout["header_row"] = ri

    if layout["ok"] is None or layout["nok"] is None:
        # Нет подписанных колонок ОК/NOK — классическая раскладка от «Области проверки».
        area = layout["area"] if layout["area"] is not None else 0
        if area + 3 >= ncols:
            return None
        layout["area"], layout["ok"], layout["nok"] = area, area + 2, area + 3
    if layout["area"] is None:
        candidates = [ci for ci in range(layout["ok"]) if ci != layout["num"]]
        if not candidates:
            return None
        layout["area"] = candidates[0]
    if layout["comment"] is None and layout["area"] + 1 < layout["ok"]:
        layout["comment"] = layout["area"] + 1
    if layout["problems"] is None and layout["nok"] + 1 < ncols:
        layout["problems"] = layout["nok"] + 1
    return layout


def _find_checklist_table(doc):
    """Первая таблица с раскладкой чек-листа: со строкой заголовков либо с титулом
    «Чек лист …» (у Чек листа Акт заголовков колонок нет)."""
    for t_idx, table in enumerate(doc.tables):
        layout = _find_layout(table)
        if not layout:
            continue
        titled = _is_title_row(table.rows[0]) and "чек" in _norm(table.rows[0].cells[0].text)
        if layout["header_row"] is not None or titled:
            return t_idx, layout
    return None, None


def parse_checklist(doc_or_path) -> list[dict]:
    """Пункты чек-листа.

    Каждый пункт: item_no, area, comments (формула / комментарий шаблона),
    problems_hint (колонка «Проблемные зоны» или последняя — источник), docs_hint,
    source_hint, ii_markers, table_index, row_index, ok_col, nok_col, problems_col.
    Строка без названия области, но с формулой — продолжение предыдущего пункта
    (Чек лист Акт: подстрока «в соответствии с п.24»), название берётся от него.
    """
    doc = _as_document(doc_or_path)
    t_idx, layout = _find_checklist_table(doc)
    if layout is None:
        return []
    table = doc.tables[t_idx]
    start = 0 if layout["header_row"] is None else layout["header_row"] + 1

    items = []
    prev_area = ""
    counter = 0
    for row_idx in range(start, len(table.rows)):
        row = table.rows[row_idx]
        cells = row.cells
        if len(cells) <= layout["nok"] or _is_title_row(row):
            continue

        def cell(key):
            ci = layout.get(key)
            return cells[ci].text.strip() if ci is not None and ci < len(cells) else ""

        num = cell("num")
        if layout["num"] is not None and not re.fullmatch(r"\d+", num):
            continue  # шапка/пустые строки в таблицах с колонкой №
        area, comment = cell("area"), cell("comment")
        if layout["comment"] == layout["area"]:
            comment = ""
        if not area and not comment:
            continue
        if not area:
            area = prev_area
        prev_area = area
        counter += 1
        problems = cell("problems")
        items.append({
            "item_no": int(num) if num else counter,
            "area": area,
            "comments": comment,
            "problems_hint": problems,
            "source_hint": problems,
            "docs_hint": cell("docs"),
            "ii_markers": re.findall(r"ИИ\d+", comment),
            "table_index": t_idx,
            "row_index": row_idx,
            "ok_col": layout["ok"],
            "nok_col": layout["nok"],
            "problems_col": layout["problems"],
        })
    return items


def _set_cell_text(cell, text: str, size: int = 10) -> None:
    for para in cell.paragraphs:
        for run in para.runs:
            run.text = ""
        para.text = ""
    if not text:
        return
    if not cell.paragraphs:
        cell.add_paragraph()
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(size)


def write_marks(doc, items: list[dict], results: list[dict], ok_mark: str, nok_mark: str,
                write_reason: bool = False, mark_size: int = 11) -> None:
    """Проставляет отметки ОК/NOK по найденным колонкам (in-place в doc).

    Пункт без вердикта (ни ok, ни nok — «проверить вручную») получает пустые клетки.
    write_reason — писать обоснование в колонку «Проблемные зоны» (Сводный акт);
    в Плане АУДИТА эта колонка — подсказка методолога, её не трогаем.
    """
    for item, res in zip(items, results):
        table = doc.tables[item["table_index"]]
        if item["row_index"] >= len(table.rows):
            continue
        cells = table.rows[item["row_index"]].cells
        is_ok = bool(res.get("ok")) and not bool(res.get("nok"))
        is_nok = bool(res.get("nok")) and not bool(res.get("ok"))
        _set_cell_text(cells[item["ok_col"]], ok_mark if is_ok else "", mark_size)
        _set_cell_text(cells[item["nok_col"]], nok_mark if is_nok else "", mark_size)
        reason = (res.get("reason") or res.get("problems") or "").strip()
        if write_reason and reason and item.get("problems_col") is not None:
            _set_cell_text(cells[item["problems_col"]], reason)


def read_marks(doc_or_path, ok_marks=("☑", "V", "✓", "✔"), nok_marks=("☒", "V", "✗", "✘")) -> list[dict]:
    """Отметки в заполненном чек-листе: [{area, ok, nok}] по колонкам из parse_checklist."""
    doc = _as_document(doc_or_path)
    rows = []
    for item in parse_checklist(doc):
        cells = doc.tables[item["table_index"]].rows[item["row_index"]].cells
        ok_text = cells[item["ok_col"]].text
        nok_text = cells[item["nok_col"]].text
        rows.append({
            "area": item["area"],
            "ok": any(m in ok_text for m in ok_marks),
            "nok": any(m in nok_text for m in nok_marks),
        })
    return rows
