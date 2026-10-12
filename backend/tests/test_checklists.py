import pytest
from docx import Document

import checklists

# (файл шаблона, тип, число пунктов) — docs/ЧЕК-ЛИСТЫ.md, §4
TEMPLATES = [
    ("ИИ -План АУДИТА.docx", "plan", 12),
    ("ЧК -Сводный акт.docx", "svod", 11),
    ("Чек лист Акт.docx", "akt", 18),  # 17 пунктов + подстрока «в соответствии с п.24»
    ("Чек лист ПН.docx", "pn", 6),
    ("Чек лист Лист оценки организацией.docx", "loo_org", 6),
    ("Чек лист Лист оценки экспертами.docx", "loo_exp", 6),
    ("Чек лист ПКД.docx", "pkd", 3),
    ("Чек лист Отчет КД.docx", "otchet_kd", 3),
    ("Чек лист Заключение.docx", "zakl", 11),
    ("Чек лист  отчет —1 этап.docx", "otchet1", 7),
    ("Чек лист Заключение —1 этап.docx", "zakl1", 7),
]


@pytest.mark.parametrize("name,ctype,count", TEMPLATES)
def test_template_type_and_items(checklist_path, name, ctype, count):
    path = checklist_path(name)
    assert checklists.detect_checklist_type(str(path)) == ctype
    items = checklists.parse_checklist(str(path))
    assert len(items) == count
    assert all(it["area"] for it in items)


def test_only_plan_and_svod_supported():
    supported = {c["key"] for c in checklists.CHECKLIST_TYPES if c["supported"]}
    assert supported == {"plan", "svod"}


def test_plan_2026_columns_found_by_headers(checklist_path):
    # 7 колонок, первая пустая: [ | Область | Комментарии | ОК | NOK | Проблемные зоны | Доки]
    items = checklists.parse_checklist(str(checklist_path("ИИ -План АУДИТА.docx")))
    first = items[0]
    assert (first["ok_col"], first["nok_col"], first["problems_col"]) == (3, 4, 5)
    assert first["area"].startswith("Дата утверждения")
    assert "ПР3" in first["comments"]
    assert first["docs_hint"].startswith("Приказ ЭГ")


def test_svod_item_numbers_from_template(checklist_path):
    items = checklists.parse_checklist(str(checklist_path("ЧК -Сводный акт.docx")))
    assert [it["item_no"] for it in items] == [1, 2, 3, 4, 5, 7, 8, 10, 11, 13, 14]
    assert items[-1]["comments"].startswith("СВ8")


def test_akt_continuation_row_inherits_area(checklist_path):
    items = checklists.parse_checklist(str(checklist_path("Чек лист Акт.docx")))
    sub = items[12]
    assert sub["comments"] == "в соответствии с п.24"
    assert sub["area"] == items[11]["area"]


def _legacy_plan(tmp_path, ok_header="ОК", nok_header="NOK"):
    """Шаблон Плана редакции 2025: 5 колонок, область в первой колонке."""
    doc = Document()
    head = doc.add_table(rows=2, cols=2)
    head.rows[0].cells[0].text = "Наименования Заявителя"
    head.rows[1].cells[0].text = "РЭГ"
    t = doc.add_table(rows=3, cols=5)
    for ci, text in enumerate(["Область проверки", "Комментарии", ok_header, nok_header, "Проблемные зоны"]):
        t.rows[0].cells[ci].text = text
    t.rows[1].cells[0].text = "Дата утверждения/согласования в шапке"
    t.rows[1].cells[1].text = "ИИ1 не позднее 5 рабочих дней до начала аудита ИИ2"
    t.rows[2].cells[0].text = "Время перерыва"
    path = tmp_path / "legacy.docx"
    doc.save(path)
    return path


@pytest.mark.parametrize("ok_header,nok_header", [("ОК", "NOK"), ("", "")])
def test_legacy_plan_layout(tmp_path, ok_header, nok_header):
    path = _legacy_plan(tmp_path, ok_header, nok_header)
    assert checklists.detect_checklist_type(str(path)) == "plan"
    items = checklists.parse_checklist(str(path))
    assert [it["area"] for it in items] == ["Дата утверждения/согласования в шапке", "Время перерыва"]
    assert (items[0]["ok_col"], items[0]["nok_col"], items[0]["problems_col"]) == (2, 3, 4)
    assert items[0]["ii_markers"] == ["ИИ1", "ИИ2"]


def test_write_and_read_marks_plan(checklist_path, tmp_path):
    path = checklist_path("ИИ -План АУДИТА.docx")
    doc = checklists.open_docx(path)
    items = checklists.parse_checklist(doc)
    results = [{"ok": i % 2 == 0, "nok": i % 2 == 1, "reason": f"r{i}"} for i in range(len(items))]
    results[3] = {"ok": False, "nok": False, "reason": "вручную"}
    checklists.write_marks(doc, items, results, ok_mark="☑", nok_mark="☒")
    out = tmp_path / "plan.docx"
    doc.save(out)

    marks = checklists.read_marks(str(out), ok_marks=("☑",), nok_marks=("☒",))
    assert [m["ok"] for m in marks] == [r["ok"] for r in results]
    assert [m["nok"] for m in marks] == [r["nok"] for r in results]
    # «Проблемные зоны» Плана — подсказки методолога, обоснование туда не пишется
    reread = checklists.parse_checklist(str(out))
    assert [it["problems_hint"] for it in reread] == [it["problems_hint"] for it in items]


def test_write_marks_svod_writes_reason(checklist_path, tmp_path):
    doc = checklists.open_docx(checklist_path("ЧК -Сводный акт.docx"))
    items = checklists.parse_checklist(doc)
    results = [{"ok": True, "nok": False, "reason": f"обоснование {it['item_no']}"} for it in items]
    checklists.write_marks(doc, items, results, ok_mark="V", nok_mark="V", write_reason=True)
    out = tmp_path / "svod.docx"
    doc.save(out)

    reread = checklists.parse_checklist(str(out))
    assert [it["problems_hint"] for it in reread] == [f"обоснование {it['item_no']}" for it in items]
    assert all(m["ok"] and not m["nok"] for m in checklists.read_marks(str(out), ("V",), ("V",)))


def test_unknown_document(tmp_path):
    doc = Document()
    doc.add_paragraph("Просто текст")
    path = tmp_path / "x.docx"
    doc.save(path)
    assert checklists.detect_checklist_type(str(path)) == "unknown"
    assert checklists.parse_checklist(str(path)) == []
